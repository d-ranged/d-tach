"""
Generate binary test fixture files (DOCX, PDF, and Excel) for manual testing.

Run once from the project root:
    python tests/fixtures/create_fixtures.py

Creates:
    tests/fixtures/sample_with_pii.docx  — DOCX with names, email, phone, student number
    tests/fixtures/sample_with_pii.pdf   — matching PDF
    tests/fixtures/sample_with_pii.xlsx  — Excel with stnum, email, name, phone columns
                                           (tests both NER mode and column-based mode)
    tests/fixtures/sample_clean.docx     — DOCX with no personal data
    tests/fixtures/sample_clean.pdf      — matching PDF
    tests/fixtures/sample_clean.xlsx     — Excel with no personal data
"""
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).parent.parent.parent))

OUTPUT_DIR = Path(__file__).parent

# ---------------------------------------------------------------------------
# DOCX fixtures
# ---------------------------------------------------------------------------

from docx import Document


def create_pii_docx() -> None:
    doc = Document()
    doc.add_heading("Internship Report — Period 2", level=1)

    info = [
        ("Student", "Sarah Mitchell"),
        ("Student Number", "542348"),
        ("Email", "sarah.mitchell@example.com"),
        ("Phone", "+44 7911 123456"),
        ("Supervisor", "Jan van der Berg"),
        ("Supervisor Email", "j.vandenberg@example.nl"),
        ("Location", "Amsterdam, Netherlands"),
        ("Date", "September 2025"),
    ]
    table = doc.add_table(rows=len(info), cols=2)
    for i, (label, value) in enumerate(info):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value

    doc.add_paragraph()
    doc.add_heading("Introduction", level=2)
    doc.add_paragraph(
        "This report was prepared by Sarah Mitchell as part of the internship programme. "
        "Student number 542348 is registered at the faculty administration."
    )
    doc.add_paragraph(
        "For questions, contact sarah.mitchell@example.com or call +44 7911 123456. "
        "The supervisor Jan van der Berg can be reached at j.vandenberg@example.nl."
    )
    doc.add_heading("Learning Objectives", level=2)
    doc.add_paragraph(
        "Sarah Mitchell identified the following learning objectives for this period. "
        "Goals were discussed with Jan van der Berg during the intake meeting in Amsterdam."
    )
    doc.add_paragraph(
        "Note: invoices totalling \u20ac12345 were processed and are not personal data. "
        "A currency-prefixed number \u20ac542348 should NOT be flagged by student number detection."
    )

    doc.save(str(OUTPUT_DIR / "sample_with_pii.docx"))
    print("Created sample_with_pii.docx")


def create_clean_docx() -> None:
    doc = Document()
    doc.add_heading("Project Notes", level=1)
    doc.add_heading("Overview", level=2)
    doc.add_paragraph(
        "This document contains general project notes with no personal information. "
        "All references are to roles, not individuals."
    )
    doc.add_heading("Timeline", level=2)
    for item in ["Q1 2025: Initial research phase", "Q2 2025: Development sprint",
                 "Q3 2025: Testing and review", "Q4 2025: Release"]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Budget", level=2)
    doc.add_paragraph("Total approved budget: \u20ac45000\nSpent to date: \u20ac12340\nRemaining: \u20ac32660")
    doc.add_heading("Notes", level=2)
    doc.add_paragraph(
        "The methodology follows standard practice. No personal data is processed or stored "
        "in this document. It should be marked as CHECKED when processed by d-tach."
    )
    doc.save(str(OUTPUT_DIR / "sample_clean.docx"))
    print("Created sample_clean.docx")


# ---------------------------------------------------------------------------
# PDF fixtures
# ---------------------------------------------------------------------------

import fitz


def create_pii_pdf() -> None:
    doc = fitz.open()
    page = doc.new_page()

    content = """\
Internship Report — Period 2

Student:          Sarah Mitchell
Student Number:   542348
Email:            sarah.mitchell@example.com
Phone:            +44 7911 123456
Supervisor:       Jan van der Berg  (j.vandenberg@example.nl)
Location:         Amsterdam, Netherlands
Date:             September 2025

Introduction

This report was prepared by Sarah Mitchell as part of the internship
programme. Student number 542348 is registered at the faculty administration.

For questions, contact sarah.mitchell@example.com or call +44 7911 123456.
The supervisor Jan van der Berg can be reached at j.vandenberg@example.nl.

Learning Objectives

Sarah Mitchell identified the following learning objectives for this period.
Goals were discussed with Jan van der Berg during the intake in Amsterdam.

Note: invoices totalling \u20ac12345 are not personal data.
A currency-prefixed number \u20ac542348 should NOT be flagged.
"""
    page.insert_text((72, 72), content, fontsize=11)
    doc.save(str(OUTPUT_DIR / "sample_with_pii.pdf"))
    doc.close()
    print("Created sample_with_pii.pdf")


def create_clean_pdf() -> None:
    doc = fitz.open()
    page = doc.new_page()
    content = """\
Project Notes

Overview

This document contains general project notes with no personal information.
All references are to roles, not individuals.

Timeline

  - Q1 2025: Initial research phase
  - Q2 2025: Development sprint
  - Q3 2025: Testing and review
  - Q4 2025: Release

Budget

  Total approved budget: \u20ac45000
  Spent to date:         \u20ac12340
  Remaining:             \u20ac32660

Notes

The methodology follows standard practice. No personal data is processed
or stored in this document. It should be marked as CHECKED by d-tach.
"""
    page.insert_text((72, 72), content, fontsize=11)
    doc.save(str(OUTPUT_DIR / "sample_clean.pdf"))
    doc.close()
    print("Created sample_clean.pdf")


# ---------------------------------------------------------------------------
# Excel fixtures
# ---------------------------------------------------------------------------

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def create_pii_xlsx() -> None:
    """Excel workbook with two sheets designed to test both processing modes.

    Sheet 1 "Students" — structured table with a header row.
      Columns include stnum (integers) and email, which the user can target
      with column-based mode. The name and phone columns contain PII that
      NER-based generic mode should detect in string cells.
      One formula cell is included to verify it is never modified.

    Sheet 2 "Supervisors" — second sheet to verify multi-sheet processing.
      Names and emails allow testing that both sheets are covered in one run.
    """
    wb = Workbook()

    # ---- Sheet 1: Students ----
    ws1 = wb.active
    ws1.title = "Students"

    header_font = Font(bold=True)
    header_fill = PatternFill(fill_type="solid", fgColor="D9E1F2")

    headers = ["stnum", "name", "email", "phone", "programme", "grade", "notes"]
    for col, h in enumerate(headers, start=1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    students = [
        (542348, "Sarah Mitchell",   "sarah.mitchell@example.com", "+44 7911 123456", "Software Engineering", 7.5, "On track"),
        (673291, "Jan van der Berg", "j.vandenberg@saxion.nl",      "+31 6 12345678",  "Software Engineering", 8.2, "Excellent progress"),
        (481027, "Fatima al-Hassan", "f.alhassan@example.com",      "+31 6 87654321",  "Data Science",         6.9, "Needs support"),
        (392846, "Thomas Becker",    "t.becker@example.de",         "+49 176 12345678","Cybersecurity",        9.1, "Outstanding"),
    ]
    for row_idx, row_data in enumerate(students, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws1.cell(row=row_idx, column=col_idx, value=value)

    # Formula cell — must survive anonymization untouched
    ws1.cell(row=6, column=6, value="=AVERAGE(F2:F5)")
    ws1.cell(row=6, column=7, value="Average grade (formula — do not anonymize)")

    ws1.column_dimensions["B"].width = 22
    ws1.column_dimensions["C"].width = 32
    ws1.column_dimensions["D"].width = 20
    ws1.column_dimensions["E"].width = 22
    ws1.column_dimensions["G"].width = 30

    # ---- Sheet 2: Supervisors ----
    ws2 = wb.create_sheet(title="Supervisors")

    sup_headers = ["name", "email", "department", "office"]
    for col, h in enumerate(sup_headers, start=1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill

    supervisors = [
        ("Dr. Anna Kowalski",  "a.kowalski@saxion.nl",   "Computer Science",  "D1.14"),
        ("Prof. Mark Johnson", "m.johnson@example.com",   "Data Engineering",  "B2.08"),
        ("Ine de Vries",       "i.devries@saxion.nl",     "Software Systems",  "D1.22"),
    ]
    for row_idx, row_data in enumerate(supervisors, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws2.cell(row=row_idx, column=col_idx, value=value)

    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 28

    wb.save(str(OUTPUT_DIR / "sample_with_pii.xlsx"))
    print("Created sample_with_pii.xlsx")


def create_clean_xlsx() -> None:
    """Excel workbook with no personal data.

    A project budget tracker — phases, descriptions, budget figures, and a
    total formula. No names, emails, student numbers, or phone numbers.
    Should be marked as CHECKED when processed by d-tach in any mode.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"

    header_font = Font(bold=True)
    header_fill = PatternFill(fill_type="solid", fgColor="E2EFDA")

    headers = ["phase", "description", "budget (€)", "spent (€)", "remaining (€)", "status"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill

    phases = [
        ("Research",    "Market and requirements analysis",     12500,  12500,  0,     "Complete"),
        ("Design",      "Architecture and UX design",           8000,   7200,   800,   "Complete"),
        ("Development", "Core implementation sprint",           45000,  31000,  14000, "In Progress"),
        ("Testing",     "QA, load testing and validation",      8000,   0,      8000,  "Pending"),
        ("Release",     "Deployment, docs and handover",        5000,   0,      5000,  "Pending"),
    ]
    for row_idx, row_data in enumerate(phases, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Totals row with formulas
    ws.cell(row=7, column=1, value="TOTAL").font = Font(bold=True)
    ws.cell(row=7, column=3, value="=SUM(C2:C6)").font = Font(bold=True)
    ws.cell(row=7, column=4, value="=SUM(D2:D6)").font = Font(bold=True)
    ws.cell(row=7, column=5, value="=SUM(E2:E6)").font = Font(bold=True)

    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["F"].width = 14

    wb.save(str(OUTPUT_DIR / "sample_clean.xlsx"))
    print("Created sample_clean.xlsx")


if __name__ == "__main__":
    create_pii_docx()
    create_clean_docx()
    create_pii_pdf()
    create_clean_pdf()
    create_pii_xlsx()
    create_clean_xlsx()
    print("Done. All fixture files created in tests/fixtures/")