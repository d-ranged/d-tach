"""Tests for the /ai/* Flask routes (issue #68 — AI Mode)."""

from pathlib import Path

import fitz
import pytest
from docx import Document
from flask import Flask
from openpyxl import Workbook

from app.routes.ai_routes import bp as ai_bp
from app.services.ai_session import AISessionStore
from app.services.anonymizer import Anonymizer
from app.services.file_processor import FileProcessor
from app.services.language_detector import LanguageDetector
from app.services.user_settings import UserSettings

TOKEN_HEADER = "X-D-Tach-Token"


def make_docx(path: Path, paragraphs: list[str]) -> None:
    """Create a DOCX file with the given paragraphs for use in tests."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))


def make_pdf(path: Path, pages: list[list[str]]) -> None:
    """Create a PDF where each inner list is one page's text lines (empty list = blank/image-only page)."""
    doc = fitz.open()
    for lines in pages:
        page = doc.new_page()
        y = 72
        for line in lines:
            page.insert_text((72, y), line)
            y += 20
    doc.save(str(path))
    doc.close()


def make_xlsx(path: Path, cells: dict[str, object]) -> None:
    """Create an xlsx file with the given cell address -> value mapping."""
    wb = Workbook()
    ws = wb.active
    for address, value in cells.items():
        ws[address] = value
    wb.save(str(path))


def _build_app(file_processor: FileProcessor, user_settings: UserSettings) -> Flask:
    """Build a minimal Flask app exposing only the ai blueprint, wired to real services."""
    flask_app = Flask(__name__)
    flask_app.register_blueprint(ai_bp)
    flask_app.file_processor = file_processor  # type: ignore[attr-defined]
    flask_app.user_settings = user_settings  # type: ignore[attr-defined]
    flask_app.ai_session_store = AISessionStore(  # type: ignore[attr-defined]
        timeout_minutes=user_settings.ai_session_timeout_minutes,
    )
    return flask_app


@pytest.fixture(scope="session")
def file_processor() -> FileProcessor:
    """Single FileProcessor shared across tests (spaCy models load once)."""
    return FileProcessor(
        anonymizer=Anonymizer(),
        language_detector=LanguageDetector(),
    )


@pytest.fixture
def user_settings(tmp_path: Path) -> UserSettings:
    """AI Mode enabled, with a temp dir under tmp_path so large-text tests stay hermetic."""
    settings = UserSettings(settings_path=tmp_path / "user_settings.json")
    settings.ai_mode_enabled = True
    settings.ai_temp_dir = str(tmp_path / "ai-temp")
    return settings


@pytest.fixture
def app(file_processor: FileProcessor, user_settings: UserSettings) -> Flask:
    return _build_app(file_processor, user_settings)


@pytest.fixture
def client(app: Flask):
    return app.test_client()


@pytest.fixture
def token(user_settings: UserSettings) -> str:
    return user_settings.ai_api_token


@pytest.fixture
def auth_headers(token: str) -> dict:
    return {TOKEN_HEADER: token}


class TestGuard:
    def test_status_ok_with_valid_token(self, client, auth_headers: dict) -> None:
        resp = client.get("/ai/status", headers=auth_headers)
        assert resp.status_code == 200
        assert resp.get_json() == {"enabled": True}

    def test_missing_token_returns_401(self, client) -> None:
        resp = client.get("/ai/status")
        assert resp.status_code == 401

    def test_wrong_token_returns_401(self, client) -> None:
        resp = client.get("/ai/status", headers={TOKEN_HEADER: "wrong-token"})
        assert resp.status_code == 401

    def test_ai_mode_disabled_returns_404(self, file_processor: FileProcessor, tmp_path: Path) -> None:
        settings = UserSettings(settings_path=tmp_path / "disabled_settings.json")
        disabled_app = _build_app(file_processor, settings)
        resp = disabled_app.test_client().get("/ai/status", headers={TOKEN_HEADER: "anything"})
        assert resp.status_code == 404


class TestExtract:
    def test_docx_with_pii_returns_anonymized_text_and_session(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.docx"
        make_docx(source, ["My name is John Smith and I work here."])

        resp = client.post("/ai/extract", json={"file_path": str(source)}, headers=auth_headers)

        assert resp.status_code == 200
        data = resp.get_json()
        assert data["session_id"]
        assert "John Smith" not in data["anonymized_text"]
        assert any(e["type"] == "PERSON" for e in data["entities"])

    def test_docx_without_pii_returns_clean_text_unchanged(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        source = tmp_path / "clean.docx"
        make_docx(source, ["The results showed a fifteen percent improvement."])

        resp = client.post("/ai/extract", json={"file_path": str(source)}, headers=auth_headers)

        assert resp.status_code == 200
        data = resp.get_json()
        assert data["anonymized_text"] == "The results showed a fifteen percent improvement."
        assert data["entities"] == []

    def test_markdown_with_pii_is_anonymized(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        source = tmp_path / "notes.md"
        source.write_text("Contact John Smith for details.", encoding="utf-8")

        resp = client.post("/ai/extract", json={"file_path": str(source)}, headers=auth_headers)

        assert resp.status_code == 200
        data = resp.get_json()
        assert "John Smith" not in data["anonymized_text"]

    def test_xlsx_with_pii_is_anonymized(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        source = tmp_path / "sheet.xlsx"
        make_xlsx(source, {"A1": "My name is John Smith."})

        resp = client.post("/ai/extract", json={"file_path": str(source)}, headers=auth_headers)

        assert resp.status_code == 200
        data = resp.get_json()
        assert "John Smith" not in data["anonymized_text"]

    def test_xlsx_without_pii_returns_clean(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        source = tmp_path / "numbers.xlsx"
        make_xlsx(source, {"A1": 123, "B1": 456})

        resp = client.post("/ai/extract", json={"file_path": str(source)}, headers=auth_headers)

        assert resp.status_code == 200
        assert resp.get_json()["entities"] == []

    def test_pdf_with_pii_is_anonymized(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.pdf"
        make_pdf(source, pages=[["My name is John Smith."]])

        resp = client.post("/ai/extract", json={"file_path": str(source)}, headers=auth_headers)

        assert resp.status_code == 200
        data = resp.get_json()
        assert "John Smith" not in data["anonymized_text"]
        assert data["warnings"] == []

    def test_pdf_fully_image_only_returns_422_unreadable(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        source = tmp_path / "scan.pdf"
        make_pdf(source, pages=[[]])

        resp = client.post("/ai/extract", json={"file_path": str(source)}, headers=auth_headers)

        assert resp.status_code == 422
        data = resp.get_json()
        assert any("image-only" in w for w in data["warnings"])

    def test_pdf_with_one_image_only_page_still_extracts_with_warning(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        source = tmp_path / "mixed.pdf"
        make_pdf(source, pages=[["My name is John Smith."], []])

        resp = client.post("/ai/extract", json={"file_path": str(source)}, headers=auth_headers)

        assert resp.status_code == 200
        data = resp.get_json()
        assert "John Smith" not in data["anonymized_text"]
        assert any("image-only" in w for w in data["warnings"])

    def test_missing_file_path_returns_400(self, client, auth_headers: dict) -> None:
        resp = client.post("/ai/extract", json={}, headers=auth_headers)
        assert resp.status_code == 400

    def test_nonexistent_file_returns_404(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        resp = client.post(
            "/ai/extract", json={"file_path": str(tmp_path / "missing.docx")}, headers=auth_headers
        )
        assert resp.status_code == 404

    def test_path_that_is_a_folder_returns_400(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        resp = client.post("/ai/extract", json={"file_path": str(tmp_path)}, headers=auth_headers)
        assert resp.status_code == 400

    def test_unsupported_extension_returns_400(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        source = tmp_path / "notes.txt"
        source.write_text("My name is John Smith.", encoding="utf-8")

        resp = client.post("/ai/extract", json={"file_path": str(source)}, headers=auth_headers)

        assert resp.status_code == 400

    def test_large_result_is_written_to_temp_file_instead_of_inlined(
        self, client, auth_headers: dict, user_settings: UserSettings, tmp_path: Path
    ) -> None:
        user_settings.ai_inline_text_max_chars = 5
        source = tmp_path / "report.docx"
        make_docx(source, ["My name is John Smith and I work here."])

        resp = client.post("/ai/extract", json={"file_path": str(source)}, headers=auth_headers)

        assert resp.status_code == 200
        data = resp.get_json()
        assert "anonymized_text" not in data
        temp_path = Path(data["anonymized_text_path"])
        assert temp_path.exists()
        assert "John Smith" not in temp_path.read_text(encoding="utf-8")


class TestRestore:
    def _extract(self, client, auth_headers: dict, tmp_path: Path, name: str = "report.docx") -> dict:
        source = tmp_path / name
        make_docx(source, ["My name is John Smith and I work here."])
        resp = client.post("/ai/extract", json={"file_path": str(source)}, headers=auth_headers)
        assert resp.status_code == 200
        return resp.get_json()

    def test_restore_inline_text_substitutes_real_values(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        extracted = self._extract(client, auth_headers, tmp_path)

        resp = client.post(
            "/ai/restore",
            json={"session_id": extracted["session_id"], "text": extracted["anonymized_text"]},
            headers=auth_headers,
        )

        assert resp.status_code == 200
        data = resp.get_json()
        assert "John Smith" in data["restored_text"]
        assert data["replacements_applied"] >= 1

    def test_restore_with_output_path_writes_to_disk_and_omits_text(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        extracted = self._extract(client, auth_headers, tmp_path)
        output_path = tmp_path / "restored.txt"

        resp = client.post(
            "/ai/restore",
            json={
                "session_id": extracted["session_id"],
                "text": extracted["anonymized_text"],
                "output_path": str(output_path),
            },
            headers=auth_headers,
        )

        assert resp.status_code == 200
        data = resp.get_json()
        assert data["status"] == "written"
        assert "restored_text" not in data
        assert "restored_text_path" not in data
        assert "John Smith" in output_path.read_text(encoding="utf-8")

    def test_restore_with_text_path_reads_from_file(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        extracted = self._extract(client, auth_headers, tmp_path)
        text_path = tmp_path / "to_restore.txt"
        text_path.write_text(extracted["anonymized_text"], encoding="utf-8")

        resp = client.post(
            "/ai/restore",
            json={"session_id": extracted["session_id"], "text_path": str(text_path)},
            headers=auth_headers,
        )

        assert resp.status_code == 200
        assert "John Smith" in resp.get_json()["restored_text"]

    def test_restore_missing_session_id_returns_400(self, client, auth_headers: dict) -> None:
        resp = client.post("/ai/restore", json={"text": "hello"}, headers=auth_headers)
        assert resp.status_code == 400

    def test_restore_unknown_session_id_returns_404(self, client, auth_headers: dict) -> None:
        resp = client.post(
            "/ai/restore", json={"session_id": "does-not-exist", "text": "hello"}, headers=auth_headers
        )
        assert resp.status_code == 404

    def test_restore_missing_text_and_text_path_returns_400(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        extracted = self._extract(client, auth_headers, tmp_path)
        resp = client.post(
            "/ai/restore", json={"session_id": extracted["session_id"]}, headers=auth_headers
        )
        assert resp.status_code == 400

    def test_restore_missing_text_path_file_returns_404(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        extracted = self._extract(client, auth_headers, tmp_path)
        resp = client.post(
            "/ai/restore",
            json={"session_id": extracted["session_id"], "text_path": str(tmp_path / "missing.txt")},
            headers=auth_headers,
        )
        assert resp.status_code == 404

    def test_large_restored_result_is_written_to_temp_file(
        self, client, auth_headers: dict, user_settings: UserSettings, tmp_path: Path
    ) -> None:
        extracted = self._extract(client, auth_headers, tmp_path)
        user_settings.ai_inline_text_max_chars = 5

        resp = client.post(
            "/ai/restore",
            json={"session_id": extracted["session_id"], "text": extracted["anonymized_text"]},
            headers=auth_headers,
        )

        assert resp.status_code == 200
        data = resp.get_json()
        assert "restored_text" not in data
        temp_path = Path(data["restored_text_path"])
        assert "John Smith" in temp_path.read_text(encoding="utf-8")


class TestRestoreFromKeyref:
    """#77 — /ai/restore accepts a KEYREF CSV, alone or merged with a session."""

    @staticmethod
    def _write_keyref(path: Path, rows: list[tuple[str, str]]) -> Path:
        lines = ["Placeholder,Original value"] + [f"{p},{o}" for p, o in rows]
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        return path

    @staticmethod
    def _extract(client, auth_headers: dict, tmp_path: Path) -> dict:
        source = tmp_path / "report.docx"
        make_docx(source, ["My name is John Smith and I work here."])
        resp = client.post("/ai/extract", json={"file_path": str(source)}, headers=auth_headers)
        assert resp.status_code == 200
        return resp.get_json()

    def test_keyref_only_restores_without_a_session(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        keyref = self._write_keyref(
            tmp_path / "KEYREF_a.csv", [("[PERSON_1]", "Ada Lovelace"), ("[PERSON_2]", "Alan Turing")]
        )

        resp = client.post(
            "/ai/restore",
            json={"keyref_path": str(keyref), "text": "[PERSON_1] met [PERSON_2]."},
            headers=auth_headers,
        )

        assert resp.status_code == 200
        data = resp.get_json()
        assert data["restored_text"] == "Ada Lovelace met Alan Turing."
        assert data["replacements_applied"] == 2

    def test_keyref_restores_a_placeholder_the_session_does_not_know(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        extracted = self._extract(client, auth_headers, tmp_path)
        keyref = self._write_keyref(tmp_path / "KEYREF_b.csv", [("[PERSON_9]", "Grace Hopper")])

        resp = client.post(
            "/ai/restore",
            json={
                "session_id": extracted["session_id"],
                "keyref_path": str(keyref),
                "text": f"{extracted['anonymized_text']} Also [PERSON_9].",
            },
            headers=auth_headers,
        )

        restored = resp.get_json()["restored_text"]
        assert "John Smith" in restored
        assert "Grace Hopper" in restored

    def test_session_wins_when_keyref_disagrees(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        extracted = self._extract(client, auth_headers, tmp_path)
        placeholder = extracted["entities"][0]["placeholder"]
        keyref = self._write_keyref(tmp_path / "KEYREF_c.csv", [(placeholder, "Stale Name")])

        resp = client.post(
            "/ai/restore",
            json={
                "session_id": extracted["session_id"],
                "keyref_path": str(keyref),
                "text": extracted["anonymized_text"],
            },
            headers=auth_headers,
        )

        restored = resp.get_json()["restored_text"]
        assert "John Smith" in restored
        assert "Stale Name" not in restored

    def test_neither_session_nor_keyref_returns_400(self, client, auth_headers: dict) -> None:
        resp = client.post("/ai/restore", json={"text": "hello"}, headers=auth_headers)
        assert resp.status_code == 400

    def test_missing_keyref_file_returns_404(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        resp = client.post(
            "/ai/restore",
            json={"keyref_path": str(tmp_path / "nope.csv"), "text": "[PERSON_1]"},
            headers=auth_headers,
        )
        assert resp.status_code == 404
        assert "KEYREF file not found" in resp.get_json()["error"]

    def test_malformed_keyref_returns_400(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        bad = tmp_path / "KEYREF_bad.csv"
        bad.write_bytes(bytes([0xFF, 0xFE, 0x00, 0x67, 0x80]))
        resp = client.post(
            "/ai/restore",
            json={"keyref_path": str(bad), "text": "[PERSON_1]"},
            headers=auth_headers,
        )
        assert resp.status_code == 400

    def test_keyref_with_no_mappings_returns_400(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        empty = self._write_keyref(tmp_path / "KEYREF_empty.csv", [])
        resp = client.post(
            "/ai/restore",
            json={"keyref_path": str(empty), "text": "[PERSON_1]"},
            headers=auth_headers,
        )
        assert resp.status_code == 400

    def test_unknown_session_still_404_even_with_keyref(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        keyref = self._write_keyref(tmp_path / "KEYREF_d.csv", [("[PERSON_1]", "Ada")])
        resp = client.post(
            "/ai/restore",
            json={"session_id": "gone", "keyref_path": str(keyref), "text": "[PERSON_1]"},
            headers=auth_headers,
        )
        assert resp.status_code == 404

    def test_output_path_still_writes_and_omits_text_with_keyref_only(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        keyref = self._write_keyref(tmp_path / "KEYREF_e.csv", [("[PERSON_1]", "Ada Lovelace")])
        output_path = tmp_path / "out" / "final.txt"

        resp = client.post(
            "/ai/restore",
            json={
                "keyref_path": str(keyref),
                "text": "Hello [PERSON_1]",
                "output_path": str(output_path),
            },
            headers=auth_headers,
        )

        assert resp.status_code == 200
        data = resp.get_json()
        assert data["status"] == "written"
        assert "restored_text" not in data
        assert "Ada Lovelace" not in str(data)
        assert output_path.read_text(encoding="utf-8") == "Hello Ada Lovelace"

    def test_large_keyref_only_result_goes_to_a_temp_file(
        self, client, auth_headers: dict, user_settings: UserSettings, tmp_path: Path
    ) -> None:
        keyref = self._write_keyref(tmp_path / "KEYREF_f.csv", [("[PERSON_1]", "Ada Lovelace")])
        user_settings.ai_inline_text_max_chars = 5

        resp = client.post(
            "/ai/restore",
            json={"keyref_path": str(keyref), "text": "Hello [PERSON_1]"},
            headers=auth_headers,
        )

        data = resp.get_json()
        assert "restored_text" not in data
        assert "Ada Lovelace" in Path(data["restored_text_path"]).read_text(encoding="utf-8")


class TestFlagTerm:
    def test_flag_term_adds_known_value(
        self, client, auth_headers: dict, user_settings: UserSettings
    ) -> None:
        resp = client.post(
            "/ai/flag-term", json={"value": "Missed Name", "entity_type": "PERSON"}, headers=auth_headers
        )

        assert resp.status_code == 200
        values = [v["value"] for v in resp.get_json()["known_values"]]
        assert "Missed Name" in values

    def test_flag_term_missing_value_returns_400(self, client, auth_headers: dict) -> None:
        resp = client.post("/ai/flag-term", json={}, headers=auth_headers)
        assert resp.status_code == 400

    def test_flag_term_duplicate_value_not_added_twice(
        self, client, auth_headers: dict
    ) -> None:
        client.post("/ai/flag-term", json={"value": "Repeat Name"}, headers=auth_headers)
        resp = client.post("/ai/flag-term", json={"value": "repeat name"}, headers=auth_headers)

        values = [v["value"] for v in resp.get_json()["known_values"]]
        assert values.count("Repeat Name") + values.count("repeat name") == 1


class TestReportedPlaceholdersMatchTheText:
    """The entity list must name the placeholders that are actually in the output.

    With hashing on, the Anonymizer's sequential numbering and the hashed token
    that lands in the text are two different strings. Reporting the sequential
    one gave an agent a placeholder that appeared nowhere in the document it had
    just been handed — and worse, [PERSON_1] means a different person in every
    file, so an agent identifying subjects across a folder merged them. That is
    precisely the confusion hashing exists to prevent.
    """

    def _extract(self, client, auth_headers: dict, source: Path) -> dict:
        resp = client.post("/ai/extract", json={"file_path": str(source)}, headers=auth_headers)
        assert resp.status_code == 200
        return resp.get_json()

    def test_every_reported_placeholder_is_in_the_text_when_hashing_is_on(
        self, client, auth_headers: dict, user_settings: UserSettings, tmp_path: Path
    ) -> None:
        user_settings.hashing_enabled = True
        user_settings.hashing_secret = "pepper"
        source = tmp_path / "notes.md"
        source.write_text("John Smith emailed john.smith@example.com.", encoding="utf-8")

        data = self._extract(client, auth_headers, source)

        assert data["entities"]
        for entity in data["entities"]:
            assert entity["placeholder"] in data["anonymized_text"], entity

    def test_reported_person_placeholder_is_hashed_not_sequential(
        self, client, auth_headers: dict, user_settings: UserSettings, tmp_path: Path
    ) -> None:
        user_settings.hashing_enabled = True
        user_settings.hashing_secret = "pepper"
        source = tmp_path / "notes.md"
        source.write_text("John Smith missed the deadline.", encoding="utf-8")

        data = self._extract(client, auth_headers, source)
        people = [e["placeholder"] for e in data["entities"] if e["type"] == "PERSON"]

        assert people and people[0] != "[PERSON_1]"

    def test_sequential_placeholders_still_reported_when_hashing_is_off(
        self, client, auth_headers: dict, tmp_path: Path
    ) -> None:
        source = tmp_path / "notes.md"
        source.write_text("John Smith missed the deadline.", encoding="utf-8")

        data = self._extract(client, auth_headers, source)
        people = [e["placeholder"] for e in data["entities"] if e["type"] == "PERSON"]

        assert people == ["[PERSON_1]"]
        assert "[PERSON_1]" in data["anonymized_text"]

    def test_the_reported_placeholder_is_the_one_restore_maps_back(
        self, client, auth_headers: dict, user_settings: UserSettings, tmp_path: Path
    ) -> None:
        """An agent that composes output from the entity list must get a working restore."""
        user_settings.hashing_enabled = True
        user_settings.hashing_secret = "pepper"
        source = tmp_path / "notes.md"
        source.write_text("John Smith missed the deadline.", encoding="utf-8")

        data = self._extract(client, auth_headers, source)
        placeholder = next(e["placeholder"] for e in data["entities"] if e["type"] == "PERSON")

        resp = client.post(
            "/ai/restore",
            json={"session_id": data["session_id"], "text": f"Re: {placeholder}"},
            headers=auth_headers,
        )

        assert resp.status_code == 200
        assert resp.get_json()["restored_text"] == "Re: John Smith"
