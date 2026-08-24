"""Tests for DocumentProcessor service class."""

from pathlib import Path

import fitz
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.services.document_processor import DocumentProcessor


def make_docx(path: Path, paragraphs: list[str]) -> None:
    """Create a DOCX file with the given paragraphs for use in tests."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))


def make_pdf(path: Path, lines: list[str]) -> None:
    """Create a simple PDF with the given lines of text for use in tests."""
    doc = fitz.open()
    page = doc.new_page()
    y = 72
    for line in lines:
        page.insert_text((72, y), line)
        y += 20
    doc.save(str(path))
    doc.close()


@pytest.fixture
def processor() -> DocumentProcessor:
    return DocumentProcessor()


# ---------------------------------------------------------------------------
# DOCX tests
# ---------------------------------------------------------------------------


class TestLoadDocx:
    def test_extracts_paragraph_text(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        docx_path = tmp_path / "test.docx"
        make_docx(docx_path, ["Hello world.", "Second paragraph."])
        text, _ = processor.load_docx(docx_path)
        assert "Hello world." in text
        assert "Second paragraph." in text

    def test_returns_document_object(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        docx_path = tmp_path / "test.docx"
        make_docx(docx_path, ["Content here."])
        _, doc = processor.load_docx(docx_path)
        assert hasattr(doc, "paragraphs") and hasattr(doc, "save")

    def test_empty_paragraphs_excluded(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        docx_path = tmp_path / "test.docx"
        make_docx(docx_path, ["", "Non-empty paragraph.", ""])
        text, _ = processor.load_docx(docx_path)
        assert text.strip() == "Non-empty paragraph."


class TestSaveDocxWithReplacements:
    def test_replacement_applied_to_output(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.docx"
        dest = tmp_path / "output.docx"
        make_docx(source, ["My name is John Smith."])
        _, doc = processor.load_docx(source)

        processor.save_docx_with_replacements(
            doc, dest, {"John Smith": "[PERSON_1]"}
        )

        result_text, _ = processor.load_docx(dest)
        assert "[PERSON_1]" in result_text
        assert "John Smith" not in result_text

    def test_original_file_not_modified(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.docx"
        dest = tmp_path / "output.docx"
        make_docx(source, ["My name is John Smith."])
        _, doc = processor.load_docx(source)
        processor.save_docx_with_replacements(doc, dest, {"John Smith": "[PERSON_1]"})

        original_text, _ = processor.load_docx(source)
        assert "John Smith" in original_text

    def test_multiple_replacements_applied(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.docx"
        dest = tmp_path / "output.docx"
        make_docx(source, ["Contact John Smith at john@example.com."])
        _, doc = processor.load_docx(source)

        processor.save_docx_with_replacements(
            doc, dest, {"John Smith": "[PERSON_1]", "john@example.com": "[EMAIL_ADDRESS_1]"}
        )

        result_text, _ = processor.load_docx(dest)
        assert "[PERSON_1]" in result_text
        assert "[EMAIL_ADDRESS_1]" in result_text

    def test_output_saved_to_dest_path(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.docx"
        dest = tmp_path / "ANON_source.docx"
        make_docx(source, ["Some text."])
        _, doc = processor.load_docx(source)
        processor.save_docx_with_replacements(doc, dest, {})
        assert dest.exists()


# ---------------------------------------------------------------------------
# PDF tests
# ---------------------------------------------------------------------------


class TestLoadPdf:
    def test_extracts_text_from_pdf(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        pdf_path = tmp_path / "test.pdf"
        make_pdf(pdf_path, ["Hello world.", "Second line."])
        text, doc = processor.load_pdf(pdf_path)
        doc.close()
        assert "Hello world." in text
        assert "Second line." in text

    def test_returns_fitz_document(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        pdf_path = tmp_path / "test.pdf"
        make_pdf(pdf_path, ["Content here."])
        _, doc = processor.load_pdf(pdf_path)
        assert isinstance(doc, fitz.Document)
        doc.close()


class TestSavePdfWithReplacements:
    def test_original_text_removed(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.pdf"
        dest = tmp_path / "output.pdf"
        make_pdf(source, ["My name is John Smith."])
        _, doc = processor.load_pdf(source)
        processor.save_pdf_with_replacements(doc, dest, {"John Smith": "[PERSON_1]"})
        doc.close()

        result_doc = fitz.open(str(dest))
        result_text = result_doc[0].get_text()
        result_doc.close()
        assert "John Smith" not in result_text

    def test_placeholder_present_in_output(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.pdf"
        dest = tmp_path / "output.pdf"
        make_pdf(source, ["My name is John Smith."])
        _, doc = processor.load_pdf(source)
        processor.save_pdf_with_replacements(doc, dest, {"John Smith": "[PERSON_1]"})
        doc.close()

        result_doc = fitz.open(str(dest))
        result_text = result_doc[0].get_text()
        result_doc.close()
        assert "[PERSON_1]" in result_text

    def test_output_file_created(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.pdf"
        dest = tmp_path / "ANON_source.pdf"
        make_pdf(source, ["Some text."])
        _, doc = processor.load_pdf(source)
        processor.save_pdf_with_replacements(doc, dest, {})
        doc.close()
        assert dest.exists()

    def test_original_file_not_modified(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.pdf"
        dest = tmp_path / "output.pdf"
        make_pdf(source, ["My name is John Smith."])
        original_mtime = source.stat().st_mtime
        _, doc = processor.load_pdf(source)
        processor.save_pdf_with_replacements(doc, dest, {"John Smith": "[PERSON_1]"})
        doc.close()
        assert source.stat().st_mtime == original_mtime


# ---------------------------------------------------------------------------
# Excel tests
# ---------------------------------------------------------------------------


def make_xlsx(path: Path, cells: dict[str, object]) -> None:
    """Create an xlsx file with the given cell address → value mapping."""
    wb = Workbook()
    ws = wb.active
    for address, value in cells.items():
        ws[address] = value
    wb.save(str(path))


class TestLoadXlsx:
    def test_extracts_string_cell_text(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        xlsx_path = tmp_path / "test.xlsx"
        make_xlsx(xlsx_path, {"A1": "John Smith", "B1": "jane@example.com"})
        text, wb = processor.load_xlsx(xlsx_path)
        assert "John Smith" in text
        assert "jane@example.com" in text

    def test_formula_cells_excluded_from_text(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        xlsx_path = tmp_path / "test.xlsx"
        make_xlsx(xlsx_path, {"A1": "Some text", "B1": "=SUM(C1:C5)"})
        text, _ = processor.load_xlsx(xlsx_path)
        assert "=SUM" not in text

    def test_numeric_cells_excluded_from_text(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        xlsx_path = tmp_path / "test.xlsx"
        make_xlsx(xlsx_path, {"A1": "Name", "B1": 12345})
        text, _ = processor.load_xlsx(xlsx_path)
        assert "12345" not in text

    def test_empty_cells_excluded_from_text(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        xlsx_path = tmp_path / "test.xlsx"
        make_xlsx(xlsx_path, {"A1": "Content", "B1": None})
        text, _ = processor.load_xlsx(xlsx_path)
        assert text.strip() == "Content"

    def test_returns_workbook_object(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        xlsx_path = tmp_path / "test.xlsx"
        make_xlsx(xlsx_path, {"A1": "text"})
        _, wb = processor.load_xlsx(xlsx_path)
        assert hasattr(wb, "worksheets") and hasattr(wb, "save")


class TestSaveXlsxWithReplacements:
    def test_replacement_applied_to_string_cell(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.xlsx"
        dest = tmp_path / "output.xlsx"
        make_xlsx(source, {"A1": "John Smith"})
        _, wb = processor.load_xlsx(source)
        processor.save_xlsx_with_replacements(wb, dest, {"John Smith": "[PERSON_1]"})
        result_wb = load_workbook(str(dest))
        assert result_wb.active["A1"].value == "[PERSON_1]"

    def test_formula_cell_not_modified(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.xlsx"
        dest = tmp_path / "output.xlsx"
        make_xlsx(source, {"A1": "John Smith", "B1": "=SUM(C1:C5)"})
        _, wb = processor.load_xlsx(source)
        processor.save_xlsx_with_replacements(wb, dest, {"John Smith": "[PERSON_1]"})
        result_wb = load_workbook(str(dest))
        assert result_wb.active["B1"].value == "=SUM(C1:C5)"

    def test_numeric_cell_not_modified(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.xlsx"
        dest = tmp_path / "output.xlsx"
        make_xlsx(source, {"A1": "John Smith", "B1": 9876543})
        _, wb = processor.load_xlsx(source)
        processor.save_xlsx_with_replacements(wb, dest, {"John Smith": "[PERSON_1]"})
        result_wb = load_workbook(str(dest))
        assert result_wb.active["B1"].value == 9876543

    def test_cell_without_pii_unchanged(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.xlsx"
        dest = tmp_path / "output.xlsx"
        make_xlsx(source, {"A1": "John Smith", "B1": "No PII here"})
        _, wb = processor.load_xlsx(source)
        processor.save_xlsx_with_replacements(wb, dest, {"John Smith": "[PERSON_1]"})
        result_wb = load_workbook(str(dest))
        assert result_wb.active["B1"].value == "No PII here"

    def test_cell_font_preserved_after_replacement(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.xlsx"
        dest = tmp_path / "output.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "John Smith"
        ws["A1"].font = Font(bold=True)
        wb.save(str(source))

        _, loaded_wb = processor.load_xlsx(source)
        processor.save_xlsx_with_replacements(loaded_wb, dest, {"John Smith": "[PERSON_1]"})

        result_wb = load_workbook(str(dest))
        assert result_wb.active["A1"].font.bold

    def test_output_saved_to_dest_path(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.xlsx"
        dest = tmp_path / "ANON_source.xlsx"
        make_xlsx(source, {"A1": "text"})
        _, wb = processor.load_xlsx(source)
        processor.save_xlsx_with_replacements(wb, dest, {})
        assert dest.exists()

    def test_original_file_not_modified(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.xlsx"
        dest = tmp_path / "output.xlsx"
        make_xlsx(source, {"A1": "John Smith"})
        original_mtime = source.stat().st_mtime
        _, wb = processor.load_xlsx(source)
        processor.save_xlsx_with_replacements(wb, dest, {"John Smith": "[PERSON_1]"})
        assert source.stat().st_mtime == original_mtime

    def test_exact_replacement_replaces_numeric_cell(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.xlsx"
        dest = tmp_path / "output.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "stnum"
        ws["A2"] = 1234567
        wb.save(str(source))
        _, loaded_wb = processor.load_xlsx(source)
        processor.save_xlsx_with_replacements(
            loaded_wb, dest, {}, exact_replacements={"1234567": "[STNUM_1]"}
        )
        result_wb = load_workbook(str(dest))
        assert result_wb.active["A2"].value == "[STNUM_1]"

    def test_exact_replacement_does_not_affect_non_matching_cells(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "source.xlsx"
        dest = tmp_path / "output.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "stnum"
        ws["A2"] = 1234567
        ws["B2"] = 9999999
        wb.save(str(source))
        _, loaded_wb = processor.load_xlsx(source)
        processor.save_xlsx_with_replacements(
            loaded_wb, dest, {}, exact_replacements={"1234567": "[STNUM_1]"}
        )
        result_wb = load_workbook(str(dest))
        assert result_wb.active["B2"].value == 9999999


class TestExtractColumnReplacements:
    def test_returns_placeholder_for_column_values(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        xlsx_path = tmp_path / "data.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "stnum"
        ws["A2"] = 1234567
        ws["A3"] = 7654321
        wb.save(str(xlsx_path))
        _, loaded_wb = processor.load_xlsx(xlsx_path)
        replacements, missing = processor.extract_column_replacements(loaded_wb, ["stnum"])
        assert "1234567" in replacements
        assert replacements["1234567"].startswith("[STNUM_")
        assert missing == []

    def test_same_value_gets_same_placeholder(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        xlsx_path = tmp_path / "data.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "stnum"
        ws["A2"] = 1001
        ws["A3"] = 1001
        wb.save(str(xlsx_path))
        _, loaded_wb = processor.load_xlsx(xlsx_path)
        replacements, _ = processor.extract_column_replacements(loaded_wb, ["stnum"])
        assert len(replacements) == 1

    def test_missing_column_reported(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        xlsx_path = tmp_path / "data.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "name"
        ws["A2"] = "Alice"
        wb.save(str(xlsx_path))
        _, loaded_wb = processor.load_xlsx(xlsx_path)
        _, missing = processor.extract_column_replacements(loaded_wb, ["stnum"])
        assert "stnum" in missing

    def test_column_name_matched_case_insensitively(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        xlsx_path = tmp_path / "data.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "StNum"
        ws["A2"] = 1234567
        wb.save(str(xlsx_path))
        _, loaded_wb = processor.load_xlsx(xlsx_path)
        replacements, missing = processor.extract_column_replacements(loaded_wb, ["stnum"])
        assert "1234567" in replacements
        assert missing == []

    def test_formula_cells_in_column_skipped(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        xlsx_path = tmp_path / "data.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "stnum"
        ws["A2"] = 1234567
        ws["A3"] = "=SUM(B2:B5)"
        wb.save(str(xlsx_path))
        _, loaded_wb = processor.load_xlsx(xlsx_path)
        replacements, _ = processor.extract_column_replacements(loaded_wb, ["stnum"])
        assert "=SUM(B2:B5)" not in replacements

    def test_multiple_columns_independent_counters(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        xlsx_path = tmp_path / "data.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "stnum"
        ws["B1"] = "email"
        ws["A2"] = 1234567
        ws["B2"] = "a@b.com"
        wb.save(str(xlsx_path))
        _, loaded_wb = processor.load_xlsx(xlsx_path)
        replacements, _ = processor.extract_column_replacements(loaded_wb, ["stnum", "email"])
        assert replacements["1234567"].startswith("[STNUM_")
        assert replacements["a@b.com"].startswith("[EMAIL_")


class TestExtractColumnReplacementsWithHashing:
    def test_hashed_placeholder_format(self, processor: DocumentProcessor) -> None:
        from openpyxl import Workbook as WB
        from app.services.hash_encoder import HashEncoder
        import tempfile, pathlib
        with tempfile.TemporaryDirectory() as d:
            p = pathlib.Path(d) / "data.xlsx"
            wb = WB()
            ws = wb.active
            ws["A1"] = "stnum"
            ws["A2"] = 542348
            wb.save(str(p))
            _, loaded_wb = processor.load_xlsx(p)
            encoder = HashEncoder("test-secret")
            replacements, _ = processor.extract_column_replacements(
                loaded_wb, ["stnum"], encoder=encoder
            )
            assert "542348" in replacements
            val = replacements["542348"]
            assert val.startswith("[STNUM_")
            assert val.endswith("]")
            assert "_1]" not in val

    def test_hashed_placeholder_consistent_across_calls(self, processor: DocumentProcessor) -> None:
        from openpyxl import Workbook as WB
        from app.services.hash_encoder import HashEncoder
        import tempfile, pathlib
        encoder = HashEncoder("test-secret")
        results = []
        for _ in range(2):
            with tempfile.TemporaryDirectory() as d:
                p = pathlib.Path(d) / "data.xlsx"
                wb = WB()
                ws = wb.active
                ws["A1"] = "stnum"
                ws["A2"] = 542348
                wb.save(str(p))
                _, loaded_wb = processor.load_xlsx(p)
                reps, _ = processor.extract_column_replacements(
                    loaded_wb, ["stnum"], encoder=encoder
                )
                results.append(reps["542348"])
        assert results[0] == results[1]

    def test_no_encoder_gives_sequential(self, processor: DocumentProcessor) -> None:
        from openpyxl import Workbook as WB
        import tempfile, pathlib
        with tempfile.TemporaryDirectory() as d:
            p = pathlib.Path(d) / "data.xlsx"
            wb = WB()
            ws = wb.active
            ws["A1"] = "stnum"
            ws["A2"] = 542348
            ws["A3"] = 673291
            wb.save(str(p))
            _, loaded_wb = processor.load_xlsx(p)
            replacements, _ = processor.extract_column_replacements(
                loaded_wb, ["stnum"]
            )
            assert "[STNUM_1]" in replacements.values()
            assert "[STNUM_2]" in replacements.values()


# ---------------------------------------------------------------------------
# Class list header/column reading (issue #66)
# ---------------------------------------------------------------------------


class TestReadXlsxHeaders:
    def test_returns_string_headers_in_order(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        xlsx_path = tmp_path / "roster.xlsx"
        make_xlsx(xlsx_path, {"A1": "Name", "B1": "Student number", "C1": "Email"})
        headers = processor.read_xlsx_headers(xlsx_path)
        assert headers == ["Name", "Student number", "Email"]

    def test_blank_and_non_string_headers_excluded(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        xlsx_path = tmp_path / "roster.xlsx"
        make_xlsx(xlsx_path, {"A1": "Name", "B1": "  ", "C1": 42})
        headers = processor.read_xlsx_headers(xlsx_path)
        assert headers == ["Name"]

    def test_only_active_sheet_is_read(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        xlsx_path = tmp_path / "roster.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Name"
        other = wb.create_sheet("Other")
        other["A1"] = "Ignored"
        wb.save(str(xlsx_path))
        headers = processor.read_xlsx_headers(xlsx_path)
        assert headers == ["Name"]


class TestReadXlsxColumns:
    def test_returns_values_per_header_case_insensitive(
        self, processor: DocumentProcessor, tmp_path: Path
    ) -> None:
        xlsx_path = tmp_path / "roster.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Name"
        ws["A2"] = "Craig Bradley"
        ws["A3"] = "Nick Surname"
        wb.save(str(xlsx_path))
        columns = processor.read_xlsx_columns(xlsx_path, ["name"])
        assert columns == {"name": ["Craig Bradley", "Nick Surname"]}

    def test_missing_header_returns_empty_list(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        xlsx_path = tmp_path / "roster.xlsx"
        make_xlsx(xlsx_path, {"A1": "Name", "A2": "Craig"})
        columns = processor.read_xlsx_columns(xlsx_path, ["Email"])
        assert columns == {"Email": []}

    def test_blank_and_formula_cells_skipped(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        xlsx_path = tmp_path / "roster.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Name"
        ws["A2"] = "Craig"
        ws["A3"] = None
        ws["A4"] = "=A2"
        ws["A5"] = "  "
        ws["A6"] = "Nick"
        wb.save(str(xlsx_path))
        columns = processor.read_xlsx_columns(xlsx_path, ["Name"])
        assert columns == {"Name": ["Craig", "Nick"]}

    def test_numeric_cell_values_are_stringified(self, processor: DocumentProcessor, tmp_path: Path) -> None:
        xlsx_path = tmp_path / "roster.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Student number"
        ws["A2"] = 1234567
        wb.save(str(xlsx_path))
        columns = processor.read_xlsx_columns(xlsx_path, ["Student number"])
        assert columns == {"Student number": ["1234567"]}
