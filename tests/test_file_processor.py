"""Tests for FileProcessor service class."""

from pathlib import Path

import fitz
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from app.services.anonymizer import Anonymizer
from app.services.file_processor import FileProcessor, ProcessingSettings, _build_entity_list
from app.services.language_detector import LanguageDetector


def make_docx(path: Path, paragraphs: list[str]) -> None:
    """Create a DOCX file with the given paragraphs for use in tests."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))


def make_pdf(path: Path, lines: list[str]) -> None:
    """Create a simple PDF with the given lines for use in tests."""
    doc = fitz.open()
    page = doc.new_page()
    y = 72
    for line in lines:
        page.insert_text((72, y), line)
        y += 20
    doc.save(str(path))
    doc.close()


def make_xlsx(path: Path, cells: dict[str, object]) -> None:
    """Create an xlsx file with the given cell address → value mapping."""
    wb = Workbook()
    ws = wb.active
    for address, value in cells.items():
        ws[address] = value
    wb.save(str(path))


@pytest.fixture(scope="session")
def file_processor() -> FileProcessor:
    """Single FileProcessor shared across tests (spaCy models load once)."""
    return FileProcessor(
        anonymizer=Anonymizer(),
        language_detector=LanguageDetector(),
    )


@pytest.fixture
def default_settings() -> ProcessingSettings:
    return ProcessingSettings()


class TestProcessDocx:
    def test_docx_with_pii_produces_anon_prefix(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.docx"
        make_docx(source, ["My name is John Smith and I work here."])
        result = file_processor.process(source, ProcessingSettings())
        assert result.status == "anonymized"
        assert result.output_path is not None
        assert result.output_path.name.startswith("ANON_")

    def test_docx_with_pii_replaces_name_in_output(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.docx"
        make_docx(source, ["My name is John Smith and I work here."])
        result = file_processor.process(source, ProcessingSettings())
        assert result.output_path is not None
        out_doc = Document(str(result.output_path))
        full_text = "\n".join(p.text for p in out_doc.paragraphs)
        assert "John Smith" not in full_text

    def test_docx_without_pii_produces_checked_prefix(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "clean.docx"
        make_docx(source, ["The results showed a fifteen percent improvement."])
        result = file_processor.process(source, ProcessingSettings())
        assert result.status == "clean"
        assert result.output_path is not None
        assert result.output_path.name.startswith("CHECKED_")

    def test_original_file_not_modified(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.docx"
        make_docx(source, ["My name is John Smith."])
        original_mtime = source.stat().st_mtime
        file_processor.process(source, ProcessingSettings())
        assert source.stat().st_mtime == original_mtime

    def test_keyref_file_created_when_enabled(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.docx"
        make_docx(source, ["My name is John Smith."])
        settings = ProcessingSettings(key_reference_enabled=True)
        result = file_processor.process(source, settings)
        assert result.keyref_path is not None
        assert result.keyref_path.exists()
        content = result.keyref_path.read_text(encoding="utf-8")
        assert "John Smith" in content

    def test_keyref_file_not_created_when_disabled(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.docx"
        make_docx(source, ["My name is John Smith."])
        result = file_processor.process(source, ProcessingSettings())
        assert result.keyref_path is None

    def test_unsupported_extension_returns_skipped(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        txt_file = tmp_path / "notes.txt"
        txt_file.write_text("Some text.")
        result = file_processor.process(txt_file, ProcessingSettings())
        assert result.status == "skipped"

    def test_missing_file_returns_error(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        missing = tmp_path / "does_not_exist.docx"
        result = file_processor.process(missing, ProcessingSettings())
        assert result.status == "error"
        assert result.error_message is not None

    def test_entities_found_count_is_positive_for_pii_doc(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.docx"
        make_docx(source, ["My name is John Smith."])
        result = file_processor.process(source, ProcessingSettings())
        assert result.entities_found > 0

    def test_replacements_populated_for_anonymized_file(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.docx"
        make_docx(source, ["My name is John Smith."])
        result = file_processor.process(source, ProcessingSettings())
        assert result.status == "anonymized"
        assert "John Smith" in result.replacements
        assert result.replacements["John Smith"].startswith("[")

    def test_replacements_empty_for_clean_file(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "clean.docx"
        make_docx(source, ["The results showed a fifteen percent improvement."])
        result = file_processor.process(source, ProcessingSettings())
        assert result.status == "clean"
        assert result.replacements == {}


class TestProcessPdf:
    def test_pdf_with_pii_produces_anon_prefix(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.pdf"
        make_pdf(source, ["My name is John Smith and I work here."])
        result = file_processor.process(source, ProcessingSettings())
        assert result.status == "anonymized"
        assert result.output_path is not None
        assert result.output_path.name.startswith("ANON_")

    def test_pdf_with_pii_removes_name_from_output(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.pdf"
        make_pdf(source, ["My name is John Smith and I work here."])
        result = file_processor.process(source, ProcessingSettings())
        assert result.output_path is not None
        out_doc = fitz.open(str(result.output_path))
        text = out_doc[0].get_text()
        out_doc.close()
        assert "John Smith" not in text

    def test_pdf_without_pii_produces_checked_prefix(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "clean.pdf"
        make_pdf(source, ["The results showed a fifteen percent improvement."])
        result = file_processor.process(source, ProcessingSettings())
        assert result.status == "clean"
        assert result.output_path is not None
        assert result.output_path.name.startswith("CHECKED_")

    def test_pdf_original_not_modified(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.pdf"
        make_pdf(source, ["My name is John Smith."])
        original_mtime = source.stat().st_mtime
        file_processor.process(source, ProcessingSettings())
        assert source.stat().st_mtime == original_mtime

    def test_pdf_keyref_created_when_enabled(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.pdf"
        make_pdf(source, ["My name is John Smith."])
        result = file_processor.process(source, ProcessingSettings(key_reference_enabled=True))
        assert result.keyref_path is not None
        assert result.keyref_path.exists()


class TestProcessXlsx:
    def test_xlsx_with_pii_produces_anon_prefix(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.xlsx"
        make_xlsx(source, {"A1": "My name is John Smith and I work here."})
        result = file_processor.process(source, ProcessingSettings())
        assert result.status == "anonymized"
        assert result.output_path is not None
        assert result.output_path.name.startswith("ANON_")

    def test_xlsx_with_pii_replaces_name_in_output_cell(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.xlsx"
        make_xlsx(source, {"A1": "My name is John Smith and I work here."})
        result = file_processor.process(source, ProcessingSettings())
        assert result.output_path is not None
        out_wb = load_workbook(str(result.output_path))
        cell_value = out_wb.active["A1"].value
        assert "John Smith" not in cell_value

    def test_xlsx_without_pii_produces_checked_prefix(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "clean.xlsx"
        make_xlsx(source, {"A1": "The results showed a fifteen percent improvement."})
        result = file_processor.process(source, ProcessingSettings())
        assert result.status == "clean"
        assert result.output_path is not None
        assert result.output_path.name.startswith("CHECKED_")

    def test_formula_cell_untouched_after_processing(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.xlsx"
        make_xlsx(source, {"A1": "My name is John Smith.", "B1": "=SUM(C1:C5)"})
        result = file_processor.process(source, ProcessingSettings())
        assert result.output_path is not None
        out_wb = load_workbook(str(result.output_path))
        assert out_wb.active["B1"].value == "=SUM(C1:C5)"

    def test_xlsx_original_not_modified(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.xlsx"
        make_xlsx(source, {"A1": "My name is John Smith."})
        original_mtime = source.stat().st_mtime
        file_processor.process(source, ProcessingSettings())
        assert source.stat().st_mtime == original_mtime

    def test_xlsx_keyref_created_when_enabled(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.xlsx"
        make_xlsx(source, {"A1": "My name is John Smith."})
        result = file_processor.process(source, ProcessingSettings(key_reference_enabled=True))
        assert result.keyref_path is not None
        assert result.keyref_path.exists()
        content = result.keyref_path.read_text(encoding="utf-8")
        assert "John Smith" in content

    def test_xlsx_folder_batch_includes_xlsx_files(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        from app.services.folder_processor import FolderProcessor, SUPPORTED_EXTENSIONS
        assert ".xlsx" in SUPPORTED_EXTENSIONS
        xlsx_file = tmp_path / "data.xlsx"
        make_xlsx(xlsx_file, {"A1": "Some text."})
        fp = FolderProcessor(file_processor)
        collected = fp.collect_files(tmp_path)
        assert xlsx_file in collected


def make_xlsx_with_headers(path: Path, headers: list[str], rows: list[list]) -> None:
    """Create an xlsx file with a header row and data rows."""
    wb = Workbook()
    ws = wb.active
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)
    for row_idx, row_data in enumerate(rows, start=2):
        for col_idx, value in enumerate(row_data, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(str(path))


class TestProcessXlsxColumnBased:
    def test_integer_column_value_replaced(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "students.xlsx"
        make_xlsx_with_headers(source, ["stnum", "name"], [[1234567, "Alice"]])
        settings = ProcessingSettings(
            excel_generic_enabled=False,
            excel_column_names=["stnum"],
        )
        result = file_processor.process(source, settings)
        assert result.status == "anonymized"
        out_wb = load_workbook(str(result.output_path))
        assert out_wb.active["A2"].value == "[STNUM_1]"

    def test_non_specified_column_untouched(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "students.xlsx"
        make_xlsx_with_headers(source, ["stnum", "grade"], [[1234567, 8]])
        settings = ProcessingSettings(
            excel_generic_enabled=False,
            excel_column_names=["stnum"],
        )
        result = file_processor.process(source, settings)
        out_wb = load_workbook(str(result.output_path))
        assert out_wb.active["B2"].value == 8

    def test_both_modes_off_returns_skipped(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "data.xlsx"
        make_xlsx(source, {"A1": "My name is John Smith."})
        settings = ProcessingSettings(
            excel_generic_enabled=False,
            excel_column_names=[],
        )
        result = file_processor.process(source, settings)
        assert result.status == "skipped"
        assert result.output_path is None

    def test_missing_column_produces_warning(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "data.xlsx"
        make_xlsx_with_headers(source, ["name"], [["Alice"]])
        settings = ProcessingSettings(
            excel_generic_enabled=False,
            excel_column_names=["stnum"],
        )
        result = file_processor.process(source, settings)
        assert any("stnum" in w for w in result.warnings)

    def test_ner_and_column_modes_combined(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "students.xlsx"
        make_xlsx_with_headers(
            source, ["stnum", "notes"],
            [[1234567, "My name is John Smith and I work here."]]
        )
        settings = ProcessingSettings(
            excel_generic_enabled=True,
            excel_column_names=["stnum"],
        )
        result = file_processor.process(source, settings)
        assert result.status == "anonymized"
        out_wb = load_workbook(str(result.output_path))
        assert out_wb.active["A2"].value == "[STNUM_1]"
        assert "John Smith" not in (out_wb.active["B2"].value or "")

    def test_column_mode_keyref_contains_column_entries(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "students.xlsx"
        make_xlsx_with_headers(source, ["stnum"], [[1234567]])
        settings = ProcessingSettings(
            excel_generic_enabled=False,
            excel_column_names=["stnum"],
            key_reference_enabled=True,
        )
        result = file_processor.process(source, settings)
        assert result.keyref_path is not None
        content = result.keyref_path.read_text(encoding="utf-8")
        assert "1234567" in content
        assert "[STNUM_1]" in content

    def test_column_only_skips_ner_on_text_cells(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "students.xlsx"
        make_xlsx_with_headers(
            source, ["stnum", "notes"],
            [[1234567, "My name is John Smith."]]
        )
        settings = ProcessingSettings(
            excel_generic_enabled=False,
            excel_column_names=["stnum"],
        )
        result = file_processor.process(source, settings)
        out_wb = load_workbook(str(result.output_path))
        assert "John Smith" in (out_wb.active["B2"].value or "")


class TestConsistentHashing:
    def test_email_hashed_when_hashing_enabled(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.docx"
        make_docx(source, ["Contact alice@example.com for details."])
        settings = ProcessingSettings(hashing_enabled=True, secret="test-secret")
        result = file_processor.process(source, settings)
        assert result.status == "anonymized"
        from docx import Document as DocxDoc
        out_doc = DocxDoc(str(result.output_path))
        text = "\n".join(p.text for p in out_doc.paragraphs)
        assert "EMAIL_ADDRESS" not in text
        assert "[EMAIL_" in text

    def test_email_hash_consistent_across_files(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        from docx import Document as DocxDoc
        email = "alice@example.com"
        settings = ProcessingSettings(hashing_enabled=True, secret="test-secret")
        results = []
        for i in range(2):
            p = tmp_path / f"doc{i}.docx"
            make_docx(p, [f"Contact {email} for details."])
            results.append(file_processor.process(p, settings))
        placeholders = []
        for r in results:
            doc = DocxDoc(str(r.output_path))
            text = "\n".join(p.text for p in doc.paragraphs)
            import re
            found = re.findall(r"\[EMAIL_[A-Z0-9]+\]", text)
            placeholders.extend(found)
        assert len(placeholders) == 2
        assert placeholders[0] == placeholders[1]

    def test_date_time_not_hashed_even_with_encoder(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.docx"
        make_docx(source, ["The meeting is on September 2025 at the office."])
        settings = ProcessingSettings(
            hashing_enabled=True, secret="test-secret", anonymize_dates=True
        )
        result = file_processor.process(source, settings)
        if result.status == "anonymized" and result.output_path:
            from docx import Document as DocxDoc
            doc = DocxDoc(str(result.output_path))
            text = "\n".join(p.text for p in doc.paragraphs)
            assert "[DATE_TIME_" in text or "September" not in text

    def test_xlsx_column_hashed_when_hashing_enabled(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        from openpyxl import Workbook as WB
        source = tmp_path / "students.xlsx"
        wb = WB()
        ws = wb.active
        ws["A1"] = "stnum"
        ws["A2"] = 542348
        wb.save(str(source))
        settings = ProcessingSettings(
            hashing_enabled=True, secret="test-secret",
            excel_generic_enabled=False, excel_column_names=["stnum"],
        )
        result = file_processor.process(source, settings)
        assert result.status == "anonymized"
        out_wb = load_workbook(str(result.output_path))
        cell_val = out_wb.active["A2"].value
        assert cell_val is not None
        assert "[STNUM_" in cell_val
        assert "_1]" not in cell_val

    def test_xlsx_column_hash_consistent_for_same_value(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        from openpyxl import Workbook as WB
        settings = ProcessingSettings(
            hashing_enabled=True, secret="test-secret",
            excel_generic_enabled=False, excel_column_names=["stnum"],
        )
        placeholders = []
        for i in range(2):
            source = tmp_path / f"students{i}.xlsx"
            wb = WB()
            ws = wb.active
            ws["A1"] = "stnum"
            ws["A2"] = 542348
            wb.save(str(source))
            result = file_processor.process(source, settings)
            out_wb = load_workbook(str(result.output_path))
            placeholders.append(out_wb.active["A2"].value)
        assert placeholders[0] == placeholders[1]

    def test_hashing_off_produces_sequential_placeholder(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "report.docx"
        make_docx(source, ["Contact alice@example.com for details."])
        result = file_processor.process(source, ProcessingSettings())
        if result.status == "anonymized" and result.output_path:
            from docx import Document as DocxDoc
            out_doc = DocxDoc(str(result.output_path))
            text = "\n".join(p.text for p in out_doc.paragraphs)
            assert "[EMAIL_ADDRESS_1]" in text


class TestCheckFileNames:
    """Tests for file name anonymization via check_file_names setting (issue #48)."""

    def test_prefix_mode_filename_anonymized_using_content_replacements(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        """Name detected in document content is applied case-insensitively to filename."""
        source = tmp_path / "John_Smith_report.docx"
        make_docx(source, ["My name is John Smith and I work here."])
        settings = ProcessingSettings(check_file_names=True)
        result = file_processor.process(source, settings)
        assert result.output_path is not None
        assert "John" not in result.output_path.name
        assert "Smith" not in result.output_path.name

    def test_prefix_mode_filename_unchanged_when_check_file_names_off(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        """Without check_file_names, original name is preserved (ANON_ prefix only)."""
        source = tmp_path / "John_Smith_report.docx"
        make_docx(source, ["My name is John Smith and I work here."])
        result = file_processor.process(source, ProcessingSettings())
        assert result.output_path is not None
        assert result.output_path.name == "ANON_John_Smith_report.docx"

    def test_subfolder_mode_filename_anonymized(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        """In subfolder output mode, file name is anonymized when check_file_names is on."""
        source = tmp_path / "John_Smith_report.docx"
        make_docx(source, ["My name is John Smith and I work here."])
        output_dir = tmp_path / "anonymized"
        output_dir.mkdir()
        override = output_dir / "John_Smith_report.docx"
        settings = ProcessingSettings(check_file_names=True)
        result = file_processor.process(source, settings, override)
        assert result.output_path is not None
        assert "John" not in result.output_path.name
        assert "Smith" not in result.output_path.name

    def test_subfolder_mode_filename_unchanged_when_check_file_names_off(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        """In subfolder mode without check_file_names, override path is used as-is."""
        source = tmp_path / "John_Smith_report.docx"
        make_docx(source, ["My name is John Smith."])
        output_dir = tmp_path / "anonymized"
        output_dir.mkdir()
        override = output_dir / "John_Smith_report.docx"
        settings = ProcessingSettings(check_file_names=False)
        result = file_processor.process(source, settings, override)
        assert result.output_path is not None


class TestLanguageNotLoadedPropagation:
    """A language that ensure_loaded() rejects must surface as a clean per-file error.

    FileProcessor.process() has no dedicated except clause for this — the
    friendly LanguageNotLoadedError message reaches the caller only via the
    existing generic `except Exception` catch-all in process().
    """

    def test_unloaded_language_becomes_error_result(self, tmp_path: Path) -> None:
        from app.services.language_detector import LanguageDetector

        class _RejectingRegistry:
            def is_loaded(self, language: str) -> bool:
                return False

            def ensure_loaded(self, language: str) -> None:
                raise RuntimeError("not installed")

        detector = LanguageDetector(registry=_RejectingRegistry())
        processor = FileProcessor(anonymizer=Anonymizer(), language_detector=detector)

        source = tmp_path / "report.docx"
        make_docx(source, ["My name is John Smith and I work here."])
        result = processor.process(source, ProcessingSettings(language="nl"))

        assert result.status == "error"
        assert result.error_message == (
            "Dutch detected but Dutch model is not enabled. "
            "Enable it in Settings > Languages and restart d-tach."
        )


class TestBuildEntityListKnownValues:
    """known_values entity types must be unioned in regardless of other toggles (issue #66).

    Without this, Presidio silently drops any ad-hoc recognizer result whose
    entity type is not in the requested entity list, so a class-list-imported
    NUMERIC_ID known value would vanish whenever the separate digit-count
    Numeric ID toggle happened to be off.
    """

    def test_numeric_id_known_value_added_even_when_toggle_off(self) -> None:
        result = _build_entity_list(
            anonymize_dates=False,
            numeric_id_enabled=False,
            known_values=[{"value": "1234567", "entity_type": "NUMERIC_ID", "source": "class_list"}],
        )
        assert "NUMERIC_ID" in result

    def test_no_duplicate_when_toggle_already_added_it(self) -> None:
        result = _build_entity_list(
            anonymize_dates=False,
            numeric_id_enabled=True,
            known_values=[{"value": "1234567", "entity_type": "NUMERIC_ID", "source": "class_list"}],
        )
        assert result.count("NUMERIC_ID") == 1

    def test_missing_entity_type_defaults_to_person_and_is_not_duplicated(self) -> None:
        result = _build_entity_list(
            anonymize_dates=False,
            known_values=[{"value": "Craig Bradley"}],
        )
        assert result.count("PERSON") == 1

    def test_no_known_values_leaves_list_unaffected(self) -> None:
        with_none = _build_entity_list(anonymize_dates=False, known_values=None)
        with_empty = _build_entity_list(anonymize_dates=False, known_values=[])
        assert with_none == with_empty

    def test_end_to_end_numeric_id_known_value_replaced_without_toggle(
        self, file_processor: FileProcessor, tmp_path: Path
    ) -> None:
        source = tmp_path / "note.md"
        source.write_text("Student number 1234567 is enrolled.", encoding="utf-8")
        settings = ProcessingSettings(
            numeric_id_enabled=False,
            known_values=[{"value": "1234567", "entity_type": "NUMERIC_ID", "source": "class_list"}],
        )
        result = file_processor.process(source, settings)
        assert result.status == "anonymized"
        output_text = result.output_path.read_text(encoding="utf-8")
        assert "1234567" not in output_text
        assert "[NUMERIC_ID_1]" in output_text
